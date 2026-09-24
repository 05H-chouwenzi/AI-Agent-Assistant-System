"""Structured knowledge-base parsers.

These parsers produce typed blocks instead of one opaque string. Metadata stays
attached to every block through chunking and vector indexing.
"""
from __future__ import annotations

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any


IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}


@dataclass
class ParsedBlock:
    content: str
    content_type: str = "text"
    metadata: dict[str, Any] = field(default_factory=dict)
    image_path: str | None = None
    ocr_text: str = ""


def _decode_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def _markdown_table(headers: list[Any], rows: list[list[Any]]) -> str:
    columns = [_clean_cell(value) or f"列{i + 1}" for i, value in enumerate(headers)]
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    for row in rows:
        normalized = list(row[: len(columns)])
        normalized.extend([""] * (len(columns) - len(normalized)))
        lines.append("| " + " | ".join(_clean_cell(value) for value in normalized) + " |")
    return "\n".join(lines)


def _image_block(path: Path, filename: str, **metadata: Any) -> ParsedBlock:
    try:
        from rag.loader import load_image
        ocr_text = load_image(path)
    except Exception:
        # Vision remains the semantic parser; OCR is a helpful pre-pass only.
        ocr_text = ""
    return ParsedBlock(
        content="",
        content_type="image",
        metadata={"filename": filename, "source": filename, **metadata},
        image_path=str(path),
        ocr_text=ocr_text,
    )


def _docx_image_path(path: Path, document: Any, shape: Any, index: int) -> Path | None:
    try:
        relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        image_part = document.part.related_parts[relationship_id]
        extension = Path(str(image_part.partname)).suffix.lower().lstrip(".") or "png"
        image_path = path.with_name(f"{path.stem}_image{index}.{extension}")
        image_path.write_bytes(image_part.blob)
        return image_path
    except Exception:
        return None


def parse_image(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    path = Path(path)
    return [_image_block(path, filename or path.name, file_type=path.suffix.lower().lstrip("."))]


def parse_pdf(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    import pymupdf

    path = Path(path)
    filename = filename or path.name
    blocks: list[ParsedBlock] = []
    with pymupdf.open(path) as document:
        for page_number, page in enumerate(document.pages(), 1):
            page_metadata = {"page_number": page_number, "file_type": "pdf"}
            section_title = ""
            try:
                table_finder = page.find_tables()
                tables = table_finder.tables if table_finder else []
            except Exception:
                tables = []
            covered_rects = [table.bbox for table in tables]

            text_payload = page.get_text("dict")
            for text_block in text_payload.get("blocks", []):
                if text_block.get("type") != 0:
                    continue
                rect = text_block.get("bbox")
                if any(
                    rect and rect[0] >= area[0] and rect[1] >= area[1]
                    and rect[2] <= area[2] and rect[3] <= area[3]
                    for area in covered_rects
                ):
                    continue
                lines = []
                for line in text_block.get("lines", []):
                    line_text = "".join(span.get("text", "") for span in line.get("spans", [])).strip()
                    if line_text:
                        lines.append(line_text)
                text = "\n".join(lines).strip()
                if not text:
                    continue
                if not section_title:
                    section_title = lines[0][:100]
                blocks.append(ParsedBlock(
                    content=text,
                    content_type="text",
                    metadata={**page_metadata, "section_title": section_title},
                ))

            for table_index, table in enumerate(tables, 1):
                rows = table.extract() or []
                rows = [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]
                if not rows:
                    continue
                markdown = _markdown_table(rows[0], rows[1:])
                blocks.append(ParsedBlock(
                    content=f"【PDF 表格 {table_index}】\n{markdown}",
                    content_type="table",
                    metadata={**page_metadata, "section_title": section_title, "table_index": table_index},
                ))

            seen_images: set[int] = set()
            for image_index, image_info in enumerate(page.get_images(full=True), 1):
                xref = int(image_info[0])
                if xref in seen_images:
                    continue
                seen_images.add(xref)
                extracted = document.extract_image(xref)
                image_ext = extracted.get("ext", "png").lower()
                if image_ext == "jpeg":
                    image_ext = "jpg"
                image_path = path.with_name(f"{path.stem}_page{page_number}_{xref}.{image_ext}")
                image_path.write_bytes(extracted["image"])
                blocks.append(_image_block(
                    image_path,
                    image_path.name,
                    **{
                        **page_metadata,
                        "section_title": section_title,
                        "image_index": image_index,
                        "source_document": filename,
                    },
                ))
    return blocks


def parse_docx(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    from docx import Document
    from docx.document import Document as _Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    path = Path(path)
    filename = filename or path.name
    document = Document(str(path))
    blocks: list[ParsedBlock] = []
    heading = ""
    heading_level = 0

    def walk(parent: Any) -> None:
        nonlocal heading, heading_level
        children = parent.element.body if isinstance(parent, _Document) else parent
        for child in children:
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, parent)
                style_name = (paragraph.style.name or "").lower()
                text = paragraph.text.strip()
                heading_match = re.fullmatch(r"heading (\d+)", style_name)
                if heading_match:
                    heading_level = int(heading_match.group(1))
                    heading = text or heading
                    if text:
                        blocks.append(ParsedBlock(
                            content=f"{'#' * heading_level} {text}",
                            content_type="text",
                            metadata={
                                "heading": heading,
                                "heading_level": heading_level,
                                "section_title": text,
                                "file_type": "docx",
                            },
                        ))
                elif text:
                    prefix = "- " if "list" in style_name else ""
                    blocks.append(ParsedBlock(
                        content=f"{prefix}{text}",
                        content_type="text",
                        metadata={
                            "heading": heading or None,
                            "heading_level": heading_level or None,
                            "section_title": heading or None,
                            "file_type": "docx",
                        },
                    ))
            elif child.tag == qn("w:tbl"):
                table = Table(child, parent)
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                if rows:
                    markdown = _markdown_table(rows[0], rows[1:])
                    blocks.append(ParsedBlock(
                        content=f"【DOCX 表格】\n{markdown}",
                        content_type="table",
                        metadata={"heading": heading or None, "section_title": heading or None, "file_type": "docx"},
                    ))

    walk(document)
    for image_index, shape in enumerate(document.inline_shapes, 1):
        image_path = _docx_image_path(path, document, shape, image_index)
        if not image_path:
            continue
        blocks.append(_image_block(
            image_path,
            image_path.name,
            heading=heading or None,
            section_title=heading or None,
            image_index=image_index,
            source_document=filename,
            file_type="docx",
        ))
    return blocks


def parse_pptx(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    path = Path(path)
    filename = filename or path.name
    presentation = Presentation(str(path))
    blocks: list[ParsedBlock] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        try:
            title = (slide.shapes.title.text or "").strip()
        except Exception:
            title = ""
        slide_meta = {"slide_number": slide_number, "section_title": title, "file_type": "pptx"}
        image_index = 0
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
                if text and shape != slide.shapes.title:
                    blocks.append(ParsedBlock(content=text, content_type="text", metadata=slide_meta))
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if rows:
                    markdown = _markdown_table(rows[0], rows[1:])
                    blocks.append(ParsedBlock(
                        content=f"【Slide {slide_number} 表格】\n{markdown}",
                        content_type="table",
                        metadata={**slide_meta, "table_index": image_index + 1},
                    ))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_index += 1
                try:
                    image_path = path.with_name(f"{path.stem}_slide{slide_number}_{image_index}.{(shape.image.ext or 'png').lower()}")
                    image_path.write_bytes(shape.image.blob)
                    blocks.append(_image_block(
                        image_path,
                        image_path.name,
                        **{**slide_meta, "image_index": image_index, "source_document": filename},
                    ))
                except Exception:
                    continue
    return blocks


def _sheet_blocks(sheet_name: str, headers: list[Any], rows: list[list[Any]], file_type: str) -> list[ParsedBlock]:
    blocks = []
    for chunk_index, start in enumerate(range(0, len(rows), 50), 1):
        row_slice = rows[start:start + 50]
        markdown = _markdown_table(headers, row_slice)
        blocks.append(ParsedBlock(
            content=f"【Sheet: {sheet_name}】\n{markdown}",
            content_type="table",
            metadata={
                "sheet_name": sheet_name,
                "columns": [_clean_cell(value) for value in headers],
                "row_start": start + 1,
                "row_end": start + len(row_slice),
                "row_count": len(row_slice),
                "chunk_index": chunk_index,
                "file_type": file_type,
            },
        ))
    return blocks


def parse_xlsx(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    from openpyxl import load_workbook

    path = Path(path)
    blocks: list[ParsedBlock] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
            if not rows:
                continue
            blocks.extend(_sheet_blocks(sheet_name, rows[0], rows[1:], "xlsx"))
    finally:
        workbook.close()
    return blocks


def parse_xls(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    import xlrd

    path = Path(path)
    blocks: list[ParsedBlock] = []
    workbook = xlrd.open_workbook(str(path))
    for sheet_name in workbook.sheet_names():
        sheet = workbook.sheet_by_name(sheet_name)
        rows = []
        for row_index in range(sheet.nrows):
            row = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                row.append(value)
            rows.append(row)
        rows = [row for row in rows if any(str(value).strip() for value in row)]
        if not rows:
            continue
        blocks.extend(_sheet_blocks(sheet_name, rows[0], rows[1:], "xls"))
    return blocks


def parse_csv(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    path = Path(path)
    text = _decode_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        class _DefaultDialect(csv.excel):
            delimiter = ","
        dialect = _DefaultDialect
    reader = csv.reader(io.StringIO(text), dialect)
    raw_rows = [row for row in reader if any(value.strip() for value in row)]
    if not raw_rows:
        return []
    headers = raw_rows[0]
    rows = []
    for row in raw_rows[1:]:
        row += [""] * (len(headers) - len(row))
        typed = []
        for value in row[: len(headers)]:
            trimmed = value.strip()
            try:
                typed.append(int(trimmed))
            except ValueError:
                try:
                    typed.append(float(trimmed))
                except ValueError:
                    typed.append(trimmed)
        rows.append(typed)
    return _sheet_blocks(path.stem, headers, rows, "csv")


def parse_markdown(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    path = Path(path)
    lines = _decode_text(path).splitlines()
    blocks: list[ParsedBlock] = []
    heading_path: list[str] = []
    current: list[str] = []
    current_level = 0

    def flush() -> None:
        nonlocal current
        content = "\n".join(current).strip()
        if content:
            blocks.append(ParsedBlock(
                content=content,
                content_type="text",
                metadata={
                    "heading": " / ".join(heading_path) or None,
                    "heading_level": current_level or None,
                    "section_title": heading_path[-1] if heading_path else None,
                    "file_type": "md",
                },
            ))
        current = []

    in_code = False
    for line in lines:
        if line.lstrip().startswith("```"):
            in_code = not in_code
            current.append(line)
            continue
        heading = None if in_code else re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_path = heading_path[: level - 1] + [title]
            current_level = level
            current.append(line)
            blocks.append(ParsedBlock(
                content=line,
                content_type="text",
                metadata={
                    "heading": " / ".join(heading_path),
                    "heading_level": level,
                    "section_title": title,
                    "file_type": "md",
                },
            ))
            continue
        current.append(line)
    flush()
    return blocks


def parse_txt(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    path = Path(path)
    text = _decode_text(path)
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    return [
        ParsedBlock(content=paragraph, content_type="text", metadata={"file_type": "txt"})
        for paragraph in paragraphs
    ]


def parse_json(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    path = Path(path)
    data = json.loads(_decode_text(path))
    blocks: list[ParsedBlock] = []

    def walk(value: Any, value_path: str, depth: int = 0) -> None:
        rendered = json.dumps(value, ensure_ascii=False, indent=2, default=str)
        if (not isinstance(value, (dict, list)) and len(rendered) <= 1200) or depth >= 8:
            blocks.append(ParsedBlock(
                content=rendered,
                content_type="text",
                metadata={"json_path": value_path, "file_type": "json"},
            ))
            return
        if isinstance(value, dict):
            for key, child in value.items():
                walk(child, f"{value_path}.{key}", depth + 1)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                walk(child, f"{value_path}[{index}]", depth + 1)
        else:
            blocks.append(ParsedBlock(
                content=str(value),
                content_type="text",
                metadata={"json_path": value_path, "file_type": "json"},
            ))

    walk(data, "$")
    return blocks


def parse_xml(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    path = Path(path)
    root = ET.parse(path).getroot()
    blocks: list[ParsedBlock] = []

    def walk(node: ET.Element, node_path: str, depth: int = 0) -> None:
        attributes = {f"@{key}": value for key, value in node.attrib.items()}
        rendered = json.dumps(
            {"attributes": attributes, "text": (node.text or "").strip(), "children": len(node)},
            ensure_ascii=False,
            indent=2,
        )
        content = ET.tostring(node, encoding="unicode").strip()
        if (not list(node) and len(content) <= 1200) or depth >= 8:
            blocks.append(ParsedBlock(
                content=f"{content}\n\n节点信息：\n{rendered}",
                content_type="text",
                metadata={"xml_path": node_path, "tag": node.tag, "file_type": "xml"},
            ))
            return
        for child in node:
            walk(child, f"{node_path}/{child.tag}", depth + 1)

    children = list(root)
    if len(children) <= 12:
        walk(root, f"/{root.tag}")
    else:
        for child in children:
            walk(child, f"/{root.tag}/{child.tag}")
    return blocks


def parse_knowledge_file(path: str | Path, filename: str | None = None) -> list[ParsedBlock]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path, filename)
    if suffix == ".docx":
        return parse_docx(path, filename)
    if suffix == ".pptx":
        return parse_pptx(path, filename)
    if suffix == ".xlsx":
        return parse_xlsx(path, filename)
    if suffix == ".xls":
        return parse_xls(path, filename)
    if suffix == ".csv":
        return parse_csv(path, filename)
    if suffix in {".md", ".markdown"}:
        return parse_markdown(path, filename)
    if suffix == ".json":
        return parse_json(path, filename)
    if suffix == ".xml":
        return parse_xml(path, filename)
    if suffix == ".txt":
        return parse_txt(path, filename)
    if suffix in IMAGE_SUFFIXES:
        return parse_image(path, filename)
    raise ValueError(f"Unsupported format: {suffix}")
