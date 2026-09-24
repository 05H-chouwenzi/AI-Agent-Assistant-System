import base64
import io

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from agent.nodes.fast_router import FastRouter
from services.chat_attachments import (
    AttachmentError,
    build_chat_message,
    cleanup_attachments,
    extract_attachment_context,
    load_attachments,
    save_upload,
)


class _Upload:
    def __init__(self, filename: str, data: bytes, content_type: str):
        self.filename = filename
        self.file = io.BytesIO(data)
        self.content_type = content_type


@pytest.mark.asyncio
async def test_upload_rejects_unsupported_and_mime_mismatch():
    with pytest.raises(AttachmentError):
        await save_upload(_Upload("clip.mp4", b"0123456789abcdef", "video/mp4"), 1)
    with pytest.raises(AttachmentError):
        await save_upload(_Upload("fake.png", b"not-a-png-file!", "image/png"), 1)


@pytest.mark.asyncio
async def test_csv_is_parsed_as_structure(tmp_path):
    upload = _Upload(
        "sales.csv",
        "employee,region,sales\nAlice,East,120\nBob,West,90\n".encode(),
        "text/csv",
    )
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    try:
        context = extract_attachment_context(resolved)
        assert '"employee"' in context
        assert '"Alice"' in context
        assert '"sales"' in context
    finally:
        cleanup_attachments([resolved])


@pytest.mark.asyncio
async def test_xlsx_parser_returns_sheet_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["employee", "sales"])
    sheet.append(["Alice", 120])
    buffer = io.BytesIO()
    workbook.save(buffer)

    upload = _Upload(
        "sales.xlsx",
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    try:
        context = extract_attachment_context(resolved)
        assert "【Sheet: Sales】" in context
        assert "Alice" in context
    finally:
        cleanup_attachments([resolved])


@pytest.mark.asyncio
async def test_text_document_and_parsers(tmp_path):
    resolved_attachments = []
    upload = _Upload("notes.md", "# 标题\n\n临时附件内容".encode(), "text/markdown")
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    resolved_attachments.append(resolved)
    context = extract_attachment_context(resolved)
    assert "标题" in context
    assert "临时附件内容" in context

    upload = _Upload("data.json", '{"name":"Alice","sales":120}'.encode(), "application/json")
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    resolved_attachments.append(resolved)
    context = extract_attachment_context(resolved)
    assert '"name": "Alice"' in context
    assert '"sales": 120' in context

    upload = _Upload("data.xml", "<record><name>Alice</name></record>".encode(), "application/xml")
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    resolved_attachments.append(resolved)
    context = extract_attachment_context(resolved)
    assert "<name>Alice</name>" in context

    document = Document()
    document.add_paragraph("临时 DOCX 内容")
    buffer = io.BytesIO()
    document.save(buffer)
    upload = _Upload(
        "notes.docx", buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    resolved_attachments.append(resolved)
    context = extract_attachment_context(resolved)
    assert "临时 DOCX 内容" in context

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "PPT 标题"
    slide.placeholders[1].text = "PPT 临时内容"
    buffer = io.BytesIO()
    presentation.save(buffer)
    upload = _Upload(
        "slides.pptx", buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    resolved_attachments.append(resolved)
    context = extract_attachment_context(resolved)
    assert "【Slide 1】" in context
    assert "PPT 标题" in context
    assert "PPT 临时内容" in context
    cleanup_attachments(resolved_attachments)


@pytest.mark.asyncio
async def test_image_message_contains_image_url_content_block():
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000d49444154789c626001000000ffff03000006000557bfabd40000000049454e44ae426082"
    )
    upload = _Upload("test.png", png, "image/png")
    attachment = await save_upload(upload, 42)
    resolved = load_attachments([attachment["id"]], 42)[0]
    try:
        message = await build_chat_message("描述一下这张图片", [resolved])
    finally:
        cleanup_attachments([resolved])

    assert isinstance(message.content, list)
    assert message.content[0]["type"] == "text"
    assert message.content[0]["text"].startswith("描述一下这张图片")
    image_block = message.content[1]
    assert image_block["type"] == "image_url"
    url = image_block["image_url"]["url"]
    assert url.startswith("data:image/png;base64,")
    assert base64.b64encode(png).decode("ascii") in url


@pytest.mark.asyncio
async def test_multi_attachment_message_keeps_image_block_and_table_text():
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000d49444154789c626001000000ffff03000006000557bfabd40000000049454e44ae426082"
    )
    image_upload = _Upload("test.png", png, "image/png")
    csv_upload = _Upload(
        "sales.csv", "employee,sales\nAlice,120\n".encode(), "text/csv"
    )
    image_attachment = await save_upload(image_upload, 42)
    csv_attachment = await save_upload(csv_upload, 42)
    resolved = load_attachments([image_attachment["id"], csv_attachment["id"]], 42)
    try:
        message = await build_chat_message("结合这两个文件分析", resolved)
    finally:
        cleanup_attachments(resolved)

    assert isinstance(message.content, list)
    assert message.content[0]["type"] == "text"
    assert "sales.csv" in message.content[0]["text"]
    assert '"Alice"' in message.content[0]["text"]
    assert any(
        item["type"] == "image_url" and item["image_url"]["url"].startswith("data:image/png;base64,")
        for item in message.content
    )


def test_fast_router_bypasses_direct_tool_when_attachments_exist():
    assert FastRouter().route("你好", has_attachments=True) is None
    assert FastRouter().route("你好", has_attachments=False) is not None
