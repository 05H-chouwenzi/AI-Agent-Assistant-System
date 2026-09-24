import io
import json
from pathlib import Path

import faiss
import numpy as np
import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

from rag.chunker import build_chunks
from rag.parsers import (
    parse_csv,
    parse_docx,
    parse_image,
    parse_json,
    parse_markdown,
    parse_pdf,
    parse_pptx,
    parse_txt,
    parse_xlsx,
    parse_xml,
)
from rag.vector_store.faiss_store import search as faiss_search
from tools.rag_tool import RAGTool
from utils.file_validation import (
    MAX_KNOWLEDGE_FILE_BYTES,
    FileValidationError,
    mime_allowed,
    signature_allowed,
)
from utils.request_context import reset_current_user_id, set_current_user_id


class _Upload:
    def __init__(self, filename: str, data: bytes, content_type: str):
        self.filename = filename
        self.file = io.BytesIO(data)
        self.content_type = content_type


def _blocks_text(blocks):
    return "\n".join(block.content for block in blocks)


def test_file_validation_rejects_video_and_mismatched_signature():
    assert not mime_allowed(".mp4", "video/mp4")
    assert not signature_allowed(".png", b"not-a-png")
    with pytest.raises(FileValidationError, match="文件过大"):
        class _Large:
            def read(self, size):
                return b"x" * (MAX_KNOWLEDGE_FILE_BYTES + 1)
        from utils.file_validation import write_upload_file
        write_upload_file(_Large(), Path("unused"), MAX_KNOWLEDGE_FILE_BYTES)


def test_txt_md_json_xml_parsers_keep_structure(tmp_path):
    txt = tmp_path / "notes.txt"
    txt.write_bytes("第一段\n\n第二段".encode("gb18030"))
    assert "第二段" in _blocks_text(parse_txt(txt))

    md = tmp_path / "doc.md"
    md.write_text("# 公司制度\n\n## 请假制度\n\n### 年假\n\n年假内容\n```code```\n", encoding="utf-8")
    md_blocks = parse_markdown(md)
    assert any(block.metadata.get("section_title") == "请假制度" for block in md_blocks)
    chunks = build_chunks(md_blocks, file_id=1, filename="doc.md")
    assert any(chunk["metadata"]["heading"] == "公司制度 / 请假制度" for chunk in chunks)

    json_path = tmp_path / "data.json"
    json_path.write_text(json.dumps({
        "company": {"financials": {"revenue": 100}},
        "items": [{"name": "A"}, {"name": "B"}],
    }, ensure_ascii=False), encoding="utf-8")
    json_blocks = parse_json(json_path)
    assert "$.company.financials.revenue" in {block.metadata["json_path"] for block in json_blocks}
    assert any("A" in block.content and "$.items[0]" in block.metadata["json_path"] for block in json_blocks)

    xml_path = tmp_path / "data.xml"
    xml_path.write_text("<company><financials><revenue>100</revenue></financials></company>", encoding="utf-8")
    xml_blocks = parse_xml(xml_path)
    assert "/company/financials/revenue" in {block.metadata["xml_path"] for block in xml_blocks}


def test_docx_and_pptx_structure(tmp_path):
    document = Document()
    document.add_heading("员工手册", level=1)
    document.add_paragraph("正式员工遵守本制度。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "报销"
    table.cell(1, 1).text = "100"
    embedded_image = tmp_path / "embedded.png"
    Image.new("RGB", (8, 8), "red").save(embedded_image, format="PNG")
    document.add_picture(str(embedded_image))
    docx_path = tmp_path / "document.docx"
    document.save(docx_path)
    docx_blocks = parse_docx(docx_path)
    assert any(block.metadata.get("section_title") == "员工手册" for block in docx_blocks)
    assert any(block.content_type == "table" and "100" in block.content for block in docx_blocks)
    assert any(block.content_type == "image" and block.image_path for block in docx_blocks)
    docx_chunks = build_chunks(docx_blocks, file_id=2, filename="document.docx")
    assert all(chunk["metadata"]["file_id"] == 2 for chunk in docx_chunks)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "销售总结"
    slide.placeholders[1].text = "销售额持续增长"
    slide.shapes.add_picture(str(embedded_image), Inches(1), Inches(3))
    pptx_path = tmp_path / "slides.pptx"
    presentation.save(pptx_path)
    pptx_blocks = parse_pptx(pptx_path)
    assert any(
        block.metadata.get("slide_number") == 1
        and block.metadata.get("section_title") == "销售总结"
        for block in pptx_blocks
    )
    assert any(block.content_type == "image" and block.image_path for block in pptx_blocks)


def test_xlsx_csv_sheet_header_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "财务数据"
    sheet.append(["公司", "营收", "净利润"])
    sheet.append(["A公司", 100, 20])
    sheet.append(["B公司", 200, 30])
    xlsx_path = tmp_path / "financial.xlsx"
    workbook.save(xlsx_path)
    xlsx_blocks = parse_xlsx(xlsx_path)
    assert xlsx_blocks[0].metadata["sheet_name"] == "财务数据"
    assert "公司" in xlsx_blocks[0].content and "A公司" in xlsx_blocks[0].content

    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("employee,region,sales\nAlice,East,120\n", encoding="utf-8")
    csv_blocks = parse_csv(csv_path)
    assert csv_blocks[0].metadata["columns"] == ["employee", "region", "sales"]
    assert "Alice" in csv_blocks[0].content


def test_pdf_and_image_blocks(tmp_path, monkeypatch):
    import pymupdf

    pdf_path = tmp_path / "document.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "AI knowledge base", fontsize=18)
    document.save(pdf_path)
    document.close()
    pdf_blocks = parse_pdf(pdf_path)
    assert pdf_blocks
    assert pdf_blocks[0].metadata["page_number"] == 1

    monkeypatch.setattr("rag.loader.load_image", lambda path: "sales chart")
    image_path = tmp_path / "chart.jpg"
    Image.new("RGB", (8, 8), "red").save(image_path, format="JPEG")
    image_blocks = parse_image(image_path)
    assert image_blocks[0].content_type == "image"
    assert image_blocks[0].image_path == str(image_path)
    assert image_blocks[0].ocr_text == "sales chart"


def test_rag_tool_requires_user_context():
    result = RAGTool().execute(query="测试")
    assert result.success is False
    assert "用户上下文缺失" in result.error


def test_faiss_search_filters_by_user_id(monkeypatch):
    dimension = 8
    index = faiss.IndexFlatIP(dimension)
    vectors = np.eye(dimension, dtype=np.float32)
    faiss.normalize_L2(vectors)
    index.add(vectors[:2])
    docs = [
        {"id": 0, "content": "A secret", "source": "a.txt", "metadata": {"user_id": 1, "file_id": 11}},
        {"id": 0, "content": "B secret", "source": "b.txt", "metadata": {"user_id": 2, "file_id": 22}},
    ]
    monkeypatch.setattr("rag.vector_store.faiss_store._get_index", lambda: index)
    monkeypatch.setattr("rag.vector_store.faiss_store._load_docs", lambda: docs)

    results_a = faiss_search(vectors[1].tolist(), top_k=1, user_id=1)
    results_b = faiss_search(vectors[1].tolist(), top_k=1, user_id=2)
    assert not results_a
    assert results_b[0]["content"] == "B secret"
    assert results_b[0]["metadata"]["user_id"] == 2
