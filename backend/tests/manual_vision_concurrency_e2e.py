import io
import time
from pathlib import Path

import pymupdf
import requests
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


BASE_URL = "http://127.0.0.1:8001"
RUN_ID = str(int(time.time()))
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "uploads" / "vision_concurrency_e2e"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def make_marker_image(path: Path, marker: str):
    image = Image.new("RGB", (900, 240), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 56)
    except OSError:
        font = ImageFont.load_default()
    draw.text((50, 80), marker, fill="black", font=font)
    image.save(path, format="PNG")


def make_single_image(path: Path):
    make_marker_image(path, f"VISIONCONC1 {RUN_ID}")


def make_plain_pdf(path: Path):
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 120), f"VISIONCONCPLAIN {RUN_ID}", fontsize=24)
    document.save(path)
    document.close()


def make_multi_image_pdf(path: Path, image_dir: Path):
    document = pymupdf.open()
    for index in range(1, 4):
        marker = f"VISIONCONCPDF{index} {RUN_ID}"
        image_path = image_dir / f"multi-{index}.png"
        make_marker_image(image_path, marker)
        page = document.new_page()
        page.insert_image(pymupdf.Rect(72, 72, 812, 312), filename=str(image_path))
    document.save(path)
    document.close()


def make_docx_with_image(path: Path, image_path: Path):
    document = Document()
    document.add_heading("Vision concurrency validation", level=1)
    document.add_paragraph("This document contains one embedded image.")
    document.add_picture(str(image_path), width=Inches(6))
    document.save(path)


def make_pptx_with_image(path: Path, image_path: Path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Vision concurrency validation"
    slide.placeholders[1].text = "Embedded image slide"
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(3), width=Inches(6))
    presentation.save(path)


def register_and_login(index: int):
    username = f"visconc_{RUN_ID}_{index}"
    password = "Strong@123"
    response = requests.post(
        f"{BASE_URL}/api/users/register",
        json={"username": username, "password": password},
        timeout=30,
    )
    response.raise_for_status()
    response = requests.post(
        f"{BASE_URL}/api/users/login",
        json={"username": username, "password": password},
        timeout=30,
    )
    response.raise_for_status()
    return response.json()["token"]


def upload(token: str, path: Path):
    with path.open("rb") as fileobj:
        response = requests.post(
            f"{BASE_URL}/api/knowledge/upload",
            headers={"Authorization": f"Bearer {token}"},
            files={"file": (path.name, fileobj)},
            timeout=120,
        )
    response.raise_for_status()
    return response.json()


def wait_for_document(token: str, doc_id: int):
    started = time.perf_counter()
    last_status = None
    for _ in range(300):
        response = requests.get(
            f"{BASE_URL}/api/knowledge/docs/{doc_id}",
            headers={"Authorization": f"Bearer {token}"},
            timeout=30,
        )
        response.raise_for_status()
        payload = response.json()
        last_status = payload
        if payload["status"] in {"completed", "failed"}:
            return payload, round((time.perf_counter() - started) * 1000)
        time.sleep(1)
    raise TimeoutError(f"document {doc_id} did not finish: {last_status}")


def search(token: str, query: str):
    response = requests.post(
        f"{BASE_URL}/api/tools/rag",
        headers={"Authorization": f"Bearer {token}"},
        json={"query": query, "top_k": 10},
        timeout=60,
    )
    response.raise_for_status()
    return response.json()


def validate_case(owner_token: str, other_token: str, path: Path, marker: str, expected_type: str):
    upload_response = upload(owner_token, path)
    doc_id = upload_response["id"]
    final_status, elapsed_ms = wait_for_document(owner_token, doc_id)
    owner_result = search(owner_token, marker)
    other_result = search(other_token, marker)
    owner_data = owner_result["data"]
    other_data = other_result["data"]
    owner_rows = owner_data.get("结果", []) if isinstance(owner_data.get("结果"), list) else []
    matched = any(marker in str(row) for row in owner_rows)
    content_types = sorted({
        row.get("内容类型")
        for row in owner_rows
        if isinstance(row, dict) and row.get("内容类型")
    })
    print(
        f"CASE={path.name} doc_id={doc_id} status={final_status['status']} "
        f"processing_ms={elapsed_ms} marker={marker} matched={matched} "
        f"owner_count={owner_data.get('文档数')} other_count={other_data.get('文档数')} "
        f"content_types={content_types} error={final_status.get('error_message')}"
    )
    assert final_status["status"] == "completed", final_status
    assert matched, owner_result
    assert owner_data.get("文档数", 0) > 0
    assert other_data.get("文档数", 0) == 0, other_result
    if expected_type:
        assert expected_type in content_types, content_types
    return doc_id


def validate_existing_case(
    owner_token: str,
    other_token: str,
    doc_id: int,
    source_name: str,
    marker: str,
    expected_type: str,
):
    final_status, elapsed_ms = wait_for_document(owner_token, doc_id)
    owner_result = search(owner_token, marker)
    other_result = search(other_token, marker)
    owner_data = owner_result["data"]
    other_data = other_result["data"]
    owner_rows = owner_data.get("结果", []) if isinstance(owner_data.get("结果"), list) else []
    matched = any(
        isinstance(row, dict)
        and row.get("来源") == source_name
        and marker in str(row.get("内容", ""))
        for row in owner_rows
    )
    content_types = sorted({
        row.get("内容类型")
        for row in owner_rows
        if isinstance(row, dict) and row.get("内容类型")
    })
    print(
        f"CASE={source_name} doc_id={doc_id} status={final_status['status']} "
        f"processing_ms={elapsed_ms} marker={marker} matched={matched} "
        f"owner_count={owner_data.get('文档数')} other_count={other_data.get('文档数')} "
        f"content_types={content_types} error={final_status.get('error_message')}"
    )
    assert final_status["status"] == "completed", final_status
    assert matched, owner_result
    assert owner_data.get("文档数", 0) > 0
    assert other_data.get("文档数", 0) == 0, other_result
    if expected_type:
        assert expected_type in content_types, content_types
    return doc_id


def main():
    owner_token = register_and_login(0)
    other_token = register_and_login(1)

    image_path = OUTPUT_DIR / f"single-{RUN_ID}.png"
    plain_pdf_path = OUTPUT_DIR / f"plain-{RUN_ID}.pdf"
    multi_pdf_path = OUTPUT_DIR / f"multi-{RUN_ID}.pdf"
    docx_path = OUTPUT_DIR / f"image-{RUN_ID}.docx"
    pptx_path = OUTPUT_DIR / f"image-{RUN_ID}.pptx"

    make_single_image(image_path)
    make_plain_pdf(plain_pdf_path)
    make_multi_image_pdf(multi_pdf_path, OUTPUT_DIR)
    make_marker_image(OUTPUT_DIR / f"docx-image-{RUN_ID}.png", f"VISIONCONCDOCX {RUN_ID}")
    make_docx_with_image(docx_path, OUTPUT_DIR / f"docx-image-{RUN_ID}.png")
    make_marker_image(OUTPUT_DIR / f"pptx-image-{RUN_ID}.png", f"VISIONCONCPPTX {RUN_ID}")
    make_pptx_with_image(pptx_path, OUTPUT_DIR / f"pptx-image-{RUN_ID}.png")

    validate_case(owner_token, other_token, image_path, f"VISIONCONC1 {RUN_ID}", "image")
    validate_case(owner_token, other_token, plain_pdf_path, f"VISIONCONCPLAIN {RUN_ID}", None)
    multi_upload = upload(owner_token, multi_pdf_path)
    for index in range(1, 4):
        validate_existing_case(
            owner_token,
            other_token,
            multi_upload["id"],
            multi_pdf_path.name,
            f"VISIONCONCPDF{index} {RUN_ID}",
            "image",
        )
    validate_case(owner_token, other_token, docx_path, f"VISIONCONCDOCX {RUN_ID}", "image")
    validate_case(owner_token, other_token, pptx_path, f"VISIONCONCPPTX {RUN_ID}", "image")

    print(f"E2E_OK run_id={RUN_ID}")


if __name__ == "__main__":
    main()
